from google import genai
import os
import json
import re
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


def load_key_env(filename="key.env"):
    env_file = Path(filename)
    if not env_file.exists():
        return

    for line in env_file.read_text().splitlines():
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue

        key, value = line.split("=", 1)
        key = key.strip()
        value = value.strip().strip('"').strip("'")
        os.environ.setdefault(key, value)


load_key_env()

API_KEY = os.getenv("GEMINI_API_KEY") or os.getenv("GOOGLE_API_KEY")

client = genai.Client(api_key=API_KEY)

DEFAULT_RESULTS_FILE = Path.home() / "Desktop" / "results.xlsx"

SUMMARY_HEADERS = [
    "Timestamp",
    "Query",
    "Risk Level",
    "Risk Score",
    "Recommendation",
    "Debt Level",
    "Revenue Growth",
    "Profitability",
    "News/Event Signal",
    "Stock Analysis",
]

RAW_HEADERS = ["Timestamp", "Query", "Result"]
ANALYTICS_HEADERS = ["Metric", "Value"]

def preview_text(text, max_length=None):
    cleaned = " ".join(str(text).split())
    if not max_length or len(cleaned) <= max_length:
        return cleaned
    return cleaned[: max_length - 3].rstrip() + "..."

def style_header_row(ws, headers):
    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(color="FFFFFF", bold=True)
    for column_index, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=column_index, value=header)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center")


def ensure_headers(ws, headers):
    current = [ws.cell(row=1, column=i).value for i in range(1, len(headers) + 1)]
    if current != headers:
        for i, header in enumerate(headers, start=1):
            ws.cell(row=1, column=i, value=header)
        style_header_row(ws, headers)


def autosize_columns(ws, max_width=60):
    for column_cells in ws.columns:
        longest = 0
        column_letter = column_cells[0].column_letter
        for cell in column_cells:
            if cell.value is None:
                continue
            text = str(cell.value)
            longest = max(longest, max((len(line) for line in text.splitlines()), default=0))
        ws.column_dimensions[column_letter].width = min(max(longest + 2, 12), max_width)


def extract_label_value(text, label):
    pattern = rf"(?im)^\s*{re.escape(label)}\s*:\s*(.+?)(?=\n\s*[A-Za-z][A-Za-z /-]*\s*:|\Z)"
    match = re.search(pattern, text, flags=re.DOTALL)
    if not match:
        return ""
    return " ".join(match.group(1).strip().split())


def extract_json_block(text):
    fenced_match = re.search(r"```(?:json)?\s*(\{.*?\})\s*```", text, flags=re.DOTALL)
    if fenced_match:
        return fenced_match.group(1)

    start = text.find("{")
    end = text.rfind("}")
    if start != -1 and end != -1 and end > start:
        return text[start : end + 1]
    return ""


def extract_fields_from_json(response_text):
    block = extract_json_block(response_text)
    if not block:
        return None

    try:
        data = json.loads(block)
    except json.JSONDecodeError:
        return None

    return {
        "risk": str(data.get("risk", "")).strip(),
        "risk_score": str(data.get("risk_score", "")).strip(),
        "recommendation": str(data.get("recommendation", "")).strip(),
        "debt_level": str(data.get("debt_level", "")).strip(),
        "revenue_growth": str(data.get("revenue_growth", "")).strip(),
        "profitability": str(data.get("profitability", "")).strip(),
        "news_signal": str(data.get("news_event", data.get("news_signal", ""))).strip(),
        "stock_analysis": str(data.get("stock_analysis", "")).strip(),
    }


def extract_fields(response_text):
    json_fields = extract_fields_from_json(response_text)
    if json_fields:
        return json_fields

    risk = extract_label_value(response_text, "Risk")
    risk_score_text = extract_label_value(response_text, "Risk Score")
    recommendation = extract_label_value(response_text, "Recommendation")
    debt_level = extract_label_value(response_text, "Debt Level")
    revenue_growth = extract_label_value(response_text, "Revenue Growth")
    profitability = extract_label_value(response_text, "Profitability")
    news_signal = extract_label_value(response_text, "News/Event") or extract_label_value(response_text, "News and Events")
    stock_analysis = extract_label_value(response_text, "Stock Analysis")

    if not risk:
        lowered = response_text.lower()
        if "high" in lowered and "risk" in lowered:
            risk = "High"
        elif "medium" in lowered and "risk" in lowered:
            risk = "Medium"
        elif "low" in lowered and "risk" in lowered:
            risk = "Low"

    risk_score = ""
    if risk_score_text:
        number_match = re.search(r"\d+(?:\.\d+)?", risk_score_text)
        risk_score = number_match.group(0) if number_match else risk_score_text

    if not stock_analysis:
        risk_part = response_text.split("Risk:", 1)
        if len(risk_part) > 1:
            after_risk = risk_part[1]
            stock_part = after_risk.split("Stock Analysis:", 1)
            stock_analysis = stock_part[1].strip() if len(stock_part) > 1 else ""
            if not risk:
                risk = stock_part[0].strip()

    return {
        "risk": risk,
        "risk_score": risk_score,
        "recommendation": recommendation,
        "debt_level": debt_level,
        "revenue_growth": revenue_growth,
        "profitability": profitability,
        "news_signal": news_signal,
        "stock_analysis": stock_analysis,
    }

def chart(user_query):
    return None
    
def analysis(user_query):
    prompt = (
        "You are one of the best researchers and retrieve reliable information. "
        "You use various resources such as CNBC, Economic Times and are able to analyse like "
        "a financial analyst. "
        "Return ONLY valid JSON. Do not include markdown fences or extra text. "
        "Use this exact schema with all keys present: "
        '{"risk":"Low|Medium|High","risk_score":0,"recommendation":"Buy|Hold|Sell",'
        '"debt_level":"Low|Moderate|High","revenue_growth":"...",'
        '"profitability":"...","news_event":"...","stock_analysis":"..."}'
        f"User request: {user_query}"
    )
    response = client.models.generate_content(
        model="gemini-3-flash-preview",
        contents=prompt,
    )
    return response.text

def save_to_excel(query, result, filename=DEFAULT_RESULTS_FILE):
    file = Path(filename)
    if file.exists():
        wb = load_workbook(file)
    else:
        wb = Workbook()

    if "Summary" in wb.sheetnames:
        summary_ws = wb["Summary"]
    else:
        summary_ws = wb.create_sheet("Summary", 0)

    if "Raw Results" in wb.sheetnames:
        raw_ws = wb["Raw Results"]
    else:
        raw_ws = wb.create_sheet("Raw Results")

    if "Analytics" in wb.sheetnames:
        analytics_ws = wb["Analytics"]
    else:
        analytics_ws = wb.create_sheet("Analytics")

    if wb.active.title == "Sheet" and len(wb.sheetnames) > 1:
        wb.remove(wb["Sheet"])

    ensure_headers(summary_ws, SUMMARY_HEADERS)
    ensure_headers(raw_ws, RAW_HEADERS)
    ensure_headers(analytics_ws, ANALYTICS_HEADERS)

    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    fields = extract_fields(result)

    raw_row = raw_ws.max_row + 1
    raw_ws.append([timestamp, query, result])
    raw_ws[f"C{raw_row}"].alignment = Alignment(wrap_text=True, vertical="top")
    raw_ws.row_dimensions[raw_row].height = 120

    summary_ws.append([
        timestamp,
        query,
        preview_text(fields["risk"]),
        fields["risk_score"],
        preview_text(fields["recommendation"]),
        preview_text(fields["debt_level"]),
        preview_text(fields["revenue_growth"]),
        preview_text(fields["profitability"]),
        preview_text(fields["news_signal"]),
        preview_text(fields["stock_analysis"]),
    ])
    summary_row = summary_ws.max_row
    for column in [3, 5, 6, 7, 8, 9, 10]:
        summary_ws.cell(row=summary_row, column=column).alignment = Alignment(wrap_text=True, vertical="top")
    summary_ws.row_dimensions[summary_row].height = 84

    risk_counts = {"High": 0, "Medium": 0, "Low": 0}
    recommendation_counts = {"Buy": 0, "Hold": 0, "Sell": 0}
    total_entries = 0
    for row in summary_ws.iter_rows(min_row=2, values_only=True):
        if not row or all(cell in (None, "") for cell in row):
            continue
        total_entries += 1
        risk_value = str(row[2] or "").strip().lower()
        recommendation_value = str(row[4] or "").strip().lower()
        for key in risk_counts:
            if risk_value.startswith(key.lower()):
                risk_counts[key] += 1
                break
        for key in recommendation_counts:
            if recommendation_value.startswith(key.lower()):
                recommendation_counts[key] += 1
                break

    analytics_ws.delete_rows(2, max(0, analytics_ws.max_row - 1))
    analytics_rows = [
        ["Total Entries", total_entries],
        ["High Risk Count", risk_counts["High"]],
        ["Medium Risk Count", risk_counts["Medium"]],
        ["Low Risk Count", risk_counts["Low"]],
        ["Buy Count", recommendation_counts["Buy"]],
        ["Hold Count", recommendation_counts["Hold"]],
        ["Sell Count", recommendation_counts["Sell"]],
    ]
    for row in analytics_rows:
        analytics_ws.append(row)

    for sheet in (summary_ws, raw_ws, analytics_ws):
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

    autosize_columns(summary_ws, max_width=110)
    autosize_columns(raw_ws, max_width=110)
    autosize_columns(analytics_ws, max_width=32)

    wb.save(file)


def main():
    user_query = input("What company do you want to research").strip()
    if user_query.lower() == "no" or user_query == "":
        print("Okay, going to sleep zzzz")
        return
    res = analysis(user_query)
    print(res)
    save_to_excel(user_query, res)


if __name__ == "__main__":
    main()